from langchain_core.documents import Document
import fitz
from app.config.settings import settings
from docx import Document as DocxDocument

from langchain_community.document_loaders import (PyPDFLoader,TextLoader,Docx2txtLoader,CSVLoader,UnstructuredWordDocumentLoader,WebBaseLoader) # type: ignore
import os ,io
from PIL import Image
import pytesseract
from app.rag.gemini import GeminiVision
from pydub import AudioSegment
 
from deepgram import DeepgramClient
pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

class DocumentLoader:
    def __init__(self, file_path):
        self.file_path = file_path

    def load_text(self):
        try:
            loader = TextLoader(self.file_path)
            return loader.load()
        except Exception as e:
            return f"Error loading text file: {str(e)}"         

    # def load_pdf(self):
    #     try:
    #         loader = PyPDFLoader(self.file_path)
    #         return loader.load()
    #     except Exception as e:
    #         return f"Error loading PDF file: {str(e)}"
    

    def load_pdf(self):
        try:
            documents = []
            vision = GeminiVision()

            pdf = fitz.open(self.file_path)
            for page_num in range(len(pdf)):
                page = pdf[page_num]
                
                # Extract Page Text

                page_text = page.get_text()

                if page_text.strip():
                    documents.append(
                        Document(
                            page_content=page_text,
                            metadata={
                                "source": self.file_path,
                                "page": page_num + 1,
                                "type": "text"
                            }
                        )
                    )

                
                # Extract Images
                
                image_list = page.get_images(full=True)

                for img_index, img in enumerate(image_list):
                    try:
                        xref = img[0]
                        base_image = pdf.extract_image(xref)
                        image_bytes = base_image["image"]
                        image = Image.open(
                            io.BytesIO(image_bytes)
                        )
                        # OCR
                        ocr_text = (
                            pytesseract
                            .image_to_string(image)
                            .strip()
                        )
                        # Gemini Description
                        description = ""

                        try:
                            description = (
                                vision.describe_img(image)
                            )
                        except Exception as vision_error:
                            print(
                                f"Vision Error: "
                                f"{vision_error}"
                            )

                        combined_content = f"""
    Description:
    {description}

    OCR Text:
    {ocr_text}
    """.strip()

                        if (
                            description.strip()
                            or ocr_text.strip()
                        ):
                            documents.append(
                                Document(
                                    page_content=combined_content,
                                    metadata={
                                        "source": self.file_path,
                                        "page": page_num + 1,
                                        "image_number":
                                            img_index + 1,
                                        "type": "image"
                                    }
                                )
                            )

                    except Exception as img_error:
                        print(
                            f"Image Processing Error "
                            f"(Page {page_num + 1}): "
                            f"{img_error}"
                        )

            pdf.close()

            return documents

        except Exception as e:
            raise Exception(
                f"Error loading PDF file: {e}"
            )


    # def load_docx(self):
    #     try:
    #         loader = Docx2txtLoader(self.file_path)
    #         return loader.load()
    #     except Exception as e:
    #         return f"Error loading DOCX file: {str(e)}"



    def load_docx(self):
        try:
            documents = []

            doc = DocxDocument(self.file_path)

            # Extract Paragraph Text
        
            for para_num, para in enumerate(doc.paragraphs):
                text = para.text.strip()
                if text:
                    documents.append(
                        Document(
                            page_content=text,
                            metadata={
                                "source": self.file_path,
                                "type": "text",
                                "paragraph": para_num + 1
                            }
                        )
                    )

            # Extract Tables

            for table_num, table in enumerate(doc.tables):

                rows = []
                for row in table.rows:
                    row_data = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    rows.append(
                        " | ".join(row_data)
                    )

                table_text = "\n".join(rows)

                documents.append(
                    Document(
                        page_content=table_text,
                        metadata={
                            "source": self.file_path,
                            "type": "table",
                            "table": table_num + 1
                        }
                    )
                )

            
            # Extract Images
            
            vision = GeminiVision()

            for rel in doc.part.rels.values():

                if (
                    "image"
                    in rel.target_ref.lower()
                ):

                    try:

                        image_bytes = (
                            rel.target_part.blob
                        )

                        image = Image.open(
                            io.BytesIO(image_bytes)
                        )

                        # OCR
                        ocr_text = (
                            pytesseract
                            .image_to_string(image)
                            .strip()
                        )

                        # Gemini Description
                        description = (
                            vision.describe_img(
                                image
                            )
                        )

                        combined_content = f"""
    Description:
    {description}

    OCR Text:
    {ocr_text}
    """.strip()

                        documents.append(
                            Document(
                                page_content=
                                combined_content,
                                metadata={
                                    "source":
                                    self.file_path,
                                    "type":
                                    "image"
                                }
                            )
                        )

                    except Exception as img_error:

                        print(
                            f"Image Error: "
                            f"{img_error}"
                        )

            return documents

        except Exception as e:

            raise Exception(
                f"Error loading DOCX file: "
                f"{e}"
            )
        
    def load_csv(self):
        try:
            loader = CSVLoader(self.file_path)
            return loader.load()
        except Exception as e:
            return f"Error loading CSV file: {str(e)}"
        
    def load_image(self):
        try:
            image = Image.open(self.file_path)

            ocr_text = pytesseract.image_to_string(image)
            vision = GeminiVision()
            description = vision.describe_img(self.file_path)

            combined_content = f"""
Description: {description}

OCR Text: {ocr_text}"""
            return [Document
                    (page_content=combined_content,
                     metadata={"source":self.file_path})]
        
        except Exception as e:
            raise Exception(
                f"Error loading image file: {e}"
            )
            
    def load_audio(self):
        try:
            deepgram = DeepgramClient(api_key=settings.dg_api_key)

            # Convert MP3 to WAV for better compatibility
            audio = AudioSegment.from_mp3(self.file_path)
            wav_buffer = io.BytesIO()
            audio.export(wav_buffer, format="wav")
            audio_bytes = wav_buffer.getvalue()

            response = deepgram.listen.v1.media.transcribe_file(
                request=audio_bytes,
                model="whisper-large",
                language="pa",          
                smart_format=True,
                punctuate=True,
            )

            transcript = (
                response.results.channels[0]
                .alternatives[0]
                .transcript
            )

            print(f"Transcript: '{transcript}'")

            if not transcript.strip():
                raise Exception("Empty transcript returned for audio file")

            return [
                Document(
                    page_content=transcript,
                    metadata={
                        "source": self.file_path,
                        "type": "audio",
                        "language": "pa"
                    }
                )
            ]

        except Exception as e:
            raise Exception(f"Error loading audio file: {e}")
        
    def load_web(self, url):
        try:
            loader = WebBaseLoader(web_url=url)
            return loader.load()

        except Exception as e:
            return f"Error loading web content: {str(e)}"
        

    
     
    def load(self):
        if not os.path.exists(self.file_path):
            raise FileNotFoundError(
                f"{self.file_path} not found"
            )
        _, file_extension = os.path.splitext(self.file_path)
        file_extension = file_extension.lower()
        if file_extension == ".pdf":
            return self.load_pdf()
        
        elif file_extension == ".txt":
            return self.load_text()
        
        elif file_extension == ".docx":
            return self.load_docx()

        elif file_extension == ".csv":
            return self.load_csv()
        
        elif file_extension in [".jpg", ".jpeg", ".png"]:
            return self.load_image()
        
        elif file_extension in [
                ".mp3",
                ".wav",
                ".m4a",
                ".ogg",
                ".flac"
            ]:
            return self.load_audio()

        else:
            raise ValueError(
                f"Unsupported file format: "
                f"{file_extension}"
            )






